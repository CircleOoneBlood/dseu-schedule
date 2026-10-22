"""Convert the staff-edited Excel workbook into the CSV files the page loads.

The newest .xlsx in data/excel/ (by last commit, so any upload name works) is read;
sheets 课表 / 行程 / 课程数量 become data/schedule.csv / trip.csv / counts.csv.
Columns are matched by header name, so extra helper columns (星期) are ignored.
"""
import csv
import datetime as dt
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
EXCEL_DIR = ROOT / 'data' / 'excel'

SHEETS = {
    '课表': ('schedule.csv', ['日期', '开始', '结束', '班级', '类别', '简体', 'English', '繁體', '直播']),
    '行程': ('trip.csv', ['日期', '简体', 'English', '繁體']),
    '课程数量': ('counts.csv', ['课程（简体）', 'English', '繁體', '数量']),
}
DATE_COLS, TIME_COLS = {'日期'}, {'开始', '结束'}


def newest_workbook():
    files = [p for p in EXCEL_DIR.glob('*.xlsx') if not p.name.startswith('~$')]
    if not files:
        sys.exit(f'no .xlsx found in {EXCEL_DIR}')

    def committed(p):
        out = subprocess.run(['git', 'log', '-1', '--format=%ct', '--', str(p)],
                             cwd=ROOT, capture_output=True, text=True).stdout.strip()
        return (int(out) if out else 0, p.stat().st_mtime)
    return max(files, key=committed)


def fmt(col, v):
    if v is None:
        return ''
    if col in DATE_COLS:
        if isinstance(v, (dt.datetime, dt.date)):
            return f'{v.month}/{v.day}'
    if col in TIME_COLS:
        if isinstance(v, dt.datetime):
            v = v.time()
        if isinstance(v, dt.time):
            return f'{v.hour:02d}:{v.minute:02d}'
        if isinstance(v, float) and 0 <= v < 1:  # a time Excel left as a day fraction
            m = round(v * 24 * 60)
            return f'{m // 60:02d}:{m % 60:02d}'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def main():
    src = newest_workbook()
    print(f'reading {src.relative_to(ROOT)}')
    wb = openpyxl.load_workbook(src, data_only=True)
    for sheet, (out, cols) in SHEETS.items():
        if sheet not in wb.sheetnames:
            sys.exit(f'sheet "{sheet}" missing in {src.name}')
        rows = wb[sheet].iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else '' for h in next(rows)]
        missing = [c for c in cols if c not in head]
        if missing:
            sys.exit(f'sheet "{sheet}" is missing columns: {", ".join(missing)}')
        idx = {c: head.index(c) for c in cols}
        n = 0
        with open(ROOT / 'data' / out, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, lineterminator='\n')
            w.writerow(cols)
            for r in rows:
                vals = [fmt(c, r[idx[c]] if idx[c] < len(r) else None) for c in cols]
                if any(vals):
                    w.writerow(vals)
                    n += 1
        print(f'  {sheet} → data/{out} ({n} rows)')


if __name__ == '__main__':
    main()
